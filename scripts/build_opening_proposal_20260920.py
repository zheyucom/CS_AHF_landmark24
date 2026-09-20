#!/usr/bin/env python3
from __future__ import annotations

import json
import math
import os
import shutil
from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as DOCX_RT
from docx.shared import Cm, Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement as PptxOxmlElement
from pptx.util import Cm as PptxCm, Pt as PptxPt


ROOT = Path(__file__).resolve().parent
TEMPLATE = Path(os.environ.get(
    "PROPOSAL_TEMPLATE",
    "/workspace/inbox/开题报告_心衰新发心源性休克预测模型_占舒羽_3.25.docx",
))
OUT_DIR = ROOT / "output"
ASSET_DIR = OUT_DIR / "figure_assets"
REPORT_NAME = "开题报告_ICU失代偿性心力衰竭患者早期血流动力学恶化或死亡预测模型_占舒羽_9.20.docx"
REPORT_PATH = OUT_DIR / REPORT_NAME
PPTX_PATH = OUT_DIR / "开题报告图示_可编辑源文件_占舒羽_20260920.pptx"

TITLE = "ICU失代偿性心力衰竭患者早期血流动力学恶化或死亡的预测模型：基于MIMIC的开发与内部验证及单中心外部验证"
COVER_TITLE = "ICU失代偿性心力衰竭患者早期血流动力学恶化或死亡的预测模型：\n基于MIMIC的开发与内部验证及单中心外部验证"

FONT_CN = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"
FONT_LATIN = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

COLORS = {
    "ink": "#263442",
    "muted": "#58697A",
    "line": "#477AA7",
    "blue": "#E5EEF6",
    "blue_strong": "#BFD6E8",
    "teal": "#E2F0EC",
    "teal_line": "#4C8E83",
    "amber": "#F8EED6",
    "amber_line": "#B98122",
    "rose": "#F6E3E0",
    "rose_line": "#B85D55",
    "grey": "#F5F7F9",
    "grey_line": "#CBD5DF",
    "white": "#FFFFFF",
}


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_CN if not bold else FONT_CN, size=size)


def hex_rgb(value: str):
    value = value.lstrip("#")
    return tuple(int(value[i:i+2], 16) for i in (0, 2, 4))


def draw_centered(draw, box, text, fnt, fill, spacing=8):
    x0, y0, x1, y1 = box
    lines = text.split("\n")
    bbox = draw.multiline_textbbox((0, 0), text, font=fnt, spacing=spacing, align="center")
    w = bbox[2] - bbox[0]
    h = bbox[3] - bbox[1]
    draw.multiline_text(((x0+x1-w)/2, (y0+y1-h)/2-bbox[1]), text, font=fnt,
                        fill=fill, spacing=spacing, align="center")


def draw_arrow(draw, start, end, color, width=8, dashed=False, head=24):
    x0, y0 = start
    x1, y1 = end
    if dashed:
        length = math.hypot(x1-x0, y1-y0)
        if length:
            ux, uy = (x1-x0)/length, (y1-y0)/length
            step, dash = 28, 16
            pos = 0
            while pos < length-head:
                e = min(pos+dash, length-head)
                draw.line((x0+ux*pos, y0+uy*pos, x0+ux*e, y0+uy*e), fill=color, width=width)
                pos += step
    else:
        draw.line((x0, y0, x1, y1), fill=color, width=width)
    angle = math.atan2(y1-y0, x1-x0)
    p1 = (x1 - head*math.cos(angle) + head*0.62*math.sin(angle),
          y1 - head*math.sin(angle) - head*0.62*math.cos(angle))
    p2 = (x1 - head*math.cos(angle) - head*0.62*math.sin(angle),
          y1 - head*math.sin(angle) + head*0.62*math.cos(angle))
    draw.polygon([(x1, y1), p1, p2], fill=color)


def rounded(draw, box, fill, outline, radius=28, width=4):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def save_svg_timeline(path: Path):
    svg = f'''<svg xmlns="http://www.w3.org/2000/svg" width="180mm" height="67.5mm" viewBox="0 0 2400 900">
<rect width="2400" height="900" fill="white"/>
<style>.t{{font-family:"PingFang SC","Microsoft YaHei","WenQuanYi Zen Hei",sans-serif;fill:{COLORS['ink']}}}.m{{fill:{COLORS['muted']}}}</style>
<text class="t" x="80" y="82" font-size="54" font-weight="700">Landmark 时间轴与信息边界</text>
<text class="t m" x="80" y="132" font-size="32">用前 12 h 可得信息预测随后 48 h 风险；活着离开 index ICU 为竞争事件</text>
<line x1="310" y1="455" x2="2110" y2="455" stroke="{COLORS['line']}" stroke-width="10"/>
<polygon points="2110,455 2070,432 2070,478" fill="{COLORS['line']}"/>
<rect x="90" y="210" width="1030" height="112" rx="24" fill="{COLORS['teal']}" stroke="{COLORS['teal_line']}" stroke-width="4"/>
<text class="t" x="605" y="258" font-size="34" text-anchor="middle" font-weight="700">DHF 表型支持窗 [T0-24 h, T12)</text>
<text class="t m" x="605" y="300" font-size="27" text-anchor="middle">HF 锚点 + 独立失代偿证据 + 客观/管理支持</text>
<rect x="330" y="345" width="790" height="86" rx="22" fill="{COLORS['blue']}" stroke="{COLORS['line']}" stroke-width="4"/>
<text class="t" x="725" y="400" font-size="32" text-anchor="middle" font-weight="700">预测器窗 [T0, T12)</text>
<rect x="1120" y="345" width="880" height="86" rx="22" fill="{COLORS['amber']}" stroke="{COLORS['amber_line']}" stroke-width="4"/>
<text class="t" x="1560" y="400" font-size="32" text-anchor="middle" font-weight="700">风险窗 [T12, min(T60, 活着出 ICU))</text>
<line x1="330" y1="165" x2="330" y2="625" stroke="{COLORS['ink']}" stroke-width="6"/>
<line x1="1120" y1="165" x2="1120" y2="625" stroke="{COLORS['ink']}" stroke-width="6"/>
<line x1="2000" y1="165" x2="2000" y2="625" stroke="{COLORS['ink']}" stroke-width="6"/>
<text class="t" x="330" y="690" font-size="38" text-anchor="middle" font-weight="700">T0</text>
<text class="t m" x="330" y="735" font-size="28" text-anchor="middle">index ICU 实际入科</text>
<text class="t" x="1120" y="690" font-size="38" text-anchor="middle" font-weight="700">T12</text>
<text class="t m" x="1120" y="735" font-size="28" text-anchor="middle">仍存活且留在 index ICU</text>
<text class="t" x="2000" y="690" font-size="38" text-anchor="middle" font-weight="700">T60</text>
<text class="t m" x="2000" y="735" font-size="28" text-anchor="middle">行政随访上限</text>
<rect x="1160" y="525" width="430" height="105" rx="22" fill="{COLORS['rose']}" stroke="{COLORS['rose_line']}" stroke-width="4"/>
<text class="t" x="1375" y="568" font-size="29" text-anchor="middle" font-weight="700">目标事件</text>
<text class="t m" x="1375" y="606" font-size="25" text-anchor="middle">血流动力学恶化或 ICU 内死亡</text>
<rect x="1625" y="525" width="340" height="105" rx="22" fill="{COLORS['teal']}" stroke="{COLORS['teal_line']}" stroke-width="4" stroke-dasharray="16 12"/>
<text class="t" x="1795" y="568" font-size="29" text-anchor="middle" font-weight="700">竞争事件</text>
<text class="t m" x="1795" y="606" font-size="25" text-anchor="middle">活着离开 index ICU</text>
<text class="t m" x="80" y="850" font-size="26">注：后写文书可用于重建时间，但不得把 T12 后可见信息输入预测器。</text>
</svg>'''
    path.write_text(svg, encoding="utf-8")


def save_png_timeline(path: Path):
    im = Image.new("RGB", (2400, 900), "white")
    d = ImageDraw.Draw(im)
    d.text((80, 45), "Landmark 时间轴与信息边界", font=font(54, True), fill=hex_rgb(COLORS["ink"]))
    d.text((80, 112), "用前 12 h 可得信息预测随后 48 h 风险；活着离开 index ICU 为竞争事件",
           font=font(32), fill=hex_rgb(COLORS["muted"]))
    draw_arrow(d, (310, 455), (2110, 455), hex_rgb(COLORS["line"]), 10, False, 40)
    rounded(d, (90, 210, 1120, 322), hex_rgb(COLORS["teal"]), hex_rgb(COLORS["teal_line"]), 24, 4)
    draw_centered(d, (90, 210, 1120, 275), "DHF 表型支持窗 [T0-24 h, T12)", font(34, True), hex_rgb(COLORS["ink"]))
    draw_centered(d, (90, 267, 1120, 322), "HF 锚点 + 独立失代偿证据 + 客观/管理支持", font(27), hex_rgb(COLORS["muted"]))
    rounded(d, (330, 345, 1120, 431), hex_rgb(COLORS["blue"]), hex_rgb(COLORS["line"]), 22, 4)
    draw_centered(d, (330, 345, 1120, 431), "预测器窗 [T0, T12)", font(32, True), hex_rgb(COLORS["ink"]))
    rounded(d, (1120, 345, 2000, 431), hex_rgb(COLORS["amber"]), hex_rgb(COLORS["amber_line"]), 22, 4)
    draw_centered(d, (1120, 345, 2000, 431), "风险窗 [T12, min(T60, 活着出 ICU))", font(31, True), hex_rgb(COLORS["ink"]))
    for x in (330, 1120, 2000):
        d.line((x, 165, x, 625), fill=hex_rgb(COLORS["ink"]), width=6)
    draw_centered(d, (200, 650, 460, 710), "T0", font(38, True), hex_rgb(COLORS["ink"]))
    draw_centered(d, (150, 700, 510, 770), "index ICU 实际入科", font(28), hex_rgb(COLORS["muted"]))
    draw_centered(d, (990, 650, 1250, 710), "T12", font(38, True), hex_rgb(COLORS["ink"]))
    draw_centered(d, (900, 700, 1340, 770), "仍存活且留在 index ICU", font(28), hex_rgb(COLORS["muted"]))
    draw_centered(d, (1870, 650, 2130, 710), "T60", font(38, True), hex_rgb(COLORS["ink"]))
    draw_centered(d, (1810, 700, 2190, 770), "行政随访上限", font(28), hex_rgb(COLORS["muted"]))
    rounded(d, (1160, 525, 1590, 630), hex_rgb(COLORS["rose"]), hex_rgb(COLORS["rose_line"]), 22, 4)
    draw_centered(d, (1160, 525, 1590, 575), "目标事件", font(29, True), hex_rgb(COLORS["ink"]))
    draw_centered(d, (1160, 570, 1590, 630), "血流动力学恶化或 ICU 内死亡", font(25), hex_rgb(COLORS["muted"]))
    rounded(d, (1625, 525, 1965, 630), hex_rgb(COLORS["teal"]), hex_rgb(COLORS["teal_line"]), 22, 4)
    draw_centered(d, (1625, 525, 1965, 575), "竞争事件", font(29, True), hex_rgb(COLORS["ink"]))
    draw_centered(d, (1625, 570, 1965, 630), "活着离开 index ICU", font(25), hex_rgb(COLORS["muted"]))
    d.text((80, 825), "注：后写文书可用于重建时间，但不得把 T12 后可见信息输入预测器。",
           font=font(26), fill=hex_rgb(COLORS["muted"]))
    im.save(path, dpi=(300, 300))


def save_svg_pipeline(path: Path):
    svg = f'''<svg xmlns="http://www.w3.org/2000/svg" width="180mm" height="105mm" viewBox="0 0 2400 1400">
<rect width="2400" height="1400" fill="white"/>
<style>.t{{font-family:"PingFang SC","Microsoft YaHei","WenQuanYi Zen Hei",sans-serif;fill:{COLORS['ink']}}}.m{{fill:{COLORS['muted']}}}</style>
<text class="t" x="80" y="82" font-size="54" font-weight="700">双库预测模型技术路线</text>
<text class="t m" x="80" y="132" font-size="31">先冻结对象、时间和结局，再开发模型并进行锁模后外部验证</text>
<rect x="80" y="205" width="330" height="170" rx="30" fill="{COLORS['blue']}" stroke="{COLORS['line']}" stroke-width="5"/>
<text class="t" x="245" y="265" font-size="34" text-anchor="middle" font-weight="700">MIMIC-IV</text><text class="t m" x="245" y="315" font-size="27" text-anchor="middle">开发 + 内部验证</text><text class="t m" x="245" y="350" font-size="24" text-anchor="middle">回顾性 index ICU</text>
<rect x="80" y="930" width="330" height="170" rx="30" fill="{COLORS['teal']}" stroke="{COLORS['teal_line']}" stroke-width="5"/>
<text class="t" x="245" y="990" font-size="34" text-anchor="middle" font-weight="700">本院数据库</text><text class="t m" x="245" y="1040" font-size="27" text-anchor="middle">锁模后外部验证</text><text class="t m" x="245" y="1075" font-size="24" text-anchor="middle">echo-supported 候选分母</text>
<g font-size="28" text-anchor="middle">
<rect x="500" y="205" width="300" height="170" rx="28" fill="{COLORS['grey']}" stroke="{COLORS['grey_line']}" stroke-width="4"/><text class="t" x="650" y="260" font-weight="700">键与 ICU 时间轴</text><text class="t m" x="650" y="305">唯一 episode</text><text class="t m" x="650" y="340">T0 / T12 / T60</text>
<rect x="880" y="205" width="350" height="170" rx="28" fill="{COLORS['teal']}" stroke="{COLORS['teal_line']}" stroke-width="4"/><text class="t" x="1055" y="255" font-weight="700">DHF 多域表型</text><text class="t m" x="1055" y="300">A：HF 锚点</text><text class="t m" x="1055" y="335">B：失代偿　C：支持</text>
<rect x="1310" y="205" width="350" height="170" rx="28" fill="{COLORS['amber']}" stroke="{COLORS['amber_line']}" stroke-width="4"/><text class="t" x="1485" y="255" font-weight="700">T12 风险集</text><text class="t m" x="1485" y="300">仍存活且留在 ICU</text><text class="t m" x="1485" y="335">排除既有严重状态代理</text>
<rect x="1740" y="205" width="530" height="170" rx="28" fill="{COLORS['rose']}" stroke="{COLORS['rose_line']}" stroke-width="4"/><text class="t" x="2005" y="250" font-weight="700">预测器 + 三态结局</text><text class="t m" x="2005" y="295">[T0,T12) 无泄漏预测器</text><text class="t m" x="2005" y="332">目标事件 / 活着出 ICU / 删失</text>
</g>
<line x1="410" y1="290" x2="500" y2="290" stroke="{COLORS['line']}" stroke-width="8"/><polygon points="500,290 468,271 468,309" fill="{COLORS['line']}"/>
<line x1="800" y1="290" x2="880" y2="290" stroke="{COLORS['line']}" stroke-width="8"/><polygon points="880,290 848,271 848,309" fill="{COLORS['line']}"/>
<line x1="1230" y1="290" x2="1310" y2="290" stroke="{COLORS['line']}" stroke-width="8"/><polygon points="1310,290 1278,271 1278,309" fill="{COLORS['line']}"/>
<line x1="1660" y1="290" x2="1740" y2="290" stroke="{COLORS['line']}" stroke-width="8"/><polygon points="1740,290 1708,271 1708,309" fill="{COLORS['line']}"/>
<rect x="500" y="500" width="570" height="200" rx="32" fill="{COLORS['blue']}" stroke="{COLORS['line']}" stroke-width="5"/>
<text class="t" x="785" y="560" font-size="34" text-anchor="middle" font-weight="700">Fine–Gray 主模型</text><text class="t m" x="785" y="610" font-size="27" text-anchor="middle">外层患者级内部验证</text><text class="t m" x="785" y="652" font-size="25" text-anchor="middle">折内 MICE / 标准化 / 收缩</text>
<rect x="1220" y="500" width="570" height="200" rx="32" fill="{COLORS['amber']}" stroke="{COLORS['amber_line']}" stroke-width="5"/>
<text class="t" x="1505" y="560" font-size="34" text-anchor="middle" font-weight="700">性能与稳健性</text><text class="t m" x="1505" y="610" font-size="27" text-anchor="middle">区分度 · Brier · 校准 · DCA</text><text class="t m" x="1505" y="652" font-size="25" text-anchor="middle">1 h person-period / 预设敏感性</text>
<rect x="1920" y="500" width="350" height="200" rx="32" fill="{COLORS['teal']}" stroke="{COLORS['teal_line']}" stroke-width="5"/>
<text class="t" x="2095" y="560" font-size="34" text-anchor="middle" font-weight="700">锁定模型</text><text class="t m" x="2095" y="610" font-size="27" text-anchor="middle">系数 + 预处理</text><text class="t m" x="2095" y="652" font-size="25" text-anchor="middle">阈值与版本封存</text>
<line x1="2005" y1="375" x2="2005" y2="455" stroke="{COLORS['line']}" stroke-width="8"/><line x1="2005" y1="455" x2="785" y2="455" stroke="{COLORS['line']}" stroke-width="8"/><line x1="785" y1="455" x2="785" y2="500" stroke="{COLORS['line']}" stroke-width="8"/><polygon points="785,500 765,468 805,468" fill="{COLORS['line']}"/>
<line x1="1070" y1="600" x2="1220" y2="600" stroke="{COLORS['line']}" stroke-width="8"/><polygon points="1220,600 1188,581 1188,619" fill="{COLORS['line']}"/>
<line x1="1790" y1="600" x2="1920" y2="600" stroke="{COLORS['line']}" stroke-width="8"/><polygon points="1920,600 1888,581 1888,619" fill="{COLORS['line']}"/>
<path d="M2095 700 L2095 835 L410 835 L410 1015" fill="none" stroke="{COLORS['teal_line']}" stroke-width="8" stroke-dasharray="22 16"/><polygon points="410,1015 390,983 430,983" fill="{COLORS['teal_line']}"/>
<text class="t m" x="1250" y="815" font-size="27" text-anchor="middle">锁模后原样应用；再校准与原始外部验证分开报告</text>
<rect x="500" y="930" width="1770" height="170" rx="30" fill="{COLORS['grey']}" stroke="{COLORS['grey_line']}" stroke-width="4"/>
<text class="t" x="1385" y="985" font-size="32" text-anchor="middle" font-weight="700">外部验证与适用性评估</text><text class="t m" x="1385" y="1030" font-size="26" text-anchor="middle">同口径时间窗、结局可测性、校准与临床净获益；报告双库测量差异</text><text class="t m" x="1385" y="1070" font-size="24" text-anchor="middle">early sepsis 为预设亚组；不因外部结果回改开发定义</text>
<rect x="80" y="1180" width="2190" height="145" rx="28" fill="white" stroke="{COLORS['grey_line']}" stroke-width="4"/>
<text class="t" x="120" y="1235" font-size="29" font-weight="700">全程控制</text><text class="t m" x="315" y="1235" font-size="26">变量字典与精确语义 allowlist　·　单位/比较符号/迟到结果审计　·　患者级重抽样　·　TRIPOD+AI 与 PROBAST+AI</text>
<text class="t m" x="120" y="1285" font-size="25">所有最终人数、事件数和性能来自同一冻结运行；AI 预审核与历史模型不作为临床金标准或正式结果。</text>
</svg>'''
    path.write_text(svg, encoding="utf-8")


def save_png_pipeline(path: Path):
    im = Image.new("RGB", (2400, 1400), "white")
    d = ImageDraw.Draw(im)
    ink, muted = hex_rgb(COLORS["ink"]), hex_rgb(COLORS["muted"])
    d.text((80, 45), "双库预测模型技术路线", font=font(54, True), fill=ink)
    d.text((80, 112), "先冻结对象、时间和结局，再开发模型并进行锁模后外部验证", font=font(31), fill=muted)

    def node(box, fill, outline, lines, title_size=34, body_size=27):
        rounded(d, box, hex_rgb(fill), hex_rgb(outline), 30, 5)
        x0, y0, x1, y1 = box
        draw_centered(d, (x0, y0+22, x1, y0+78), lines[0], font(title_size, True), ink)
        if len(lines) > 1:
            draw_centered(d, (x0+8, y0+76, x1-8, y1-10), "\n".join(lines[1:]), font(body_size), muted, 8)

    node((80, 205, 410, 375), COLORS["blue"], COLORS["line"], ["MIMIC-IV", "开发 + 内部验证", "回顾性 index ICU"])
    node((80, 930, 410, 1100), COLORS["teal"], COLORS["teal_line"], ["本院数据库", "锁模后外部验证", "echo-supported 候选分母"])
    node((500, 205, 800, 375), COLORS["grey"], COLORS["grey_line"], ["键与 ICU 时间轴", "唯一 episode", "T0 / T12 / T60"], 30, 25)
    node((880, 205, 1230, 375), COLORS["teal"], COLORS["teal_line"], ["DHF 多域表型", "A：HF 锚点", "B：失代偿　C：支持"], 30, 25)
    node((1310, 205, 1660, 375), COLORS["amber"], COLORS["amber_line"], ["T12 风险集", "仍存活且留在 ICU", "排除既有严重状态代理"], 30, 25)
    node((1740, 205, 2270, 375), COLORS["rose"], COLORS["rose_line"], ["预测器 + 三态结局", "[T0,T12) 无泄漏预测器", "目标事件 / 活着出 ICU / 删失"], 30, 25)
    for start, end in [((410,290),(500,290)),((800,290),(880,290)),((1230,290),(1310,290)),((1660,290),(1740,290))]:
        draw_arrow(d, start, end, hex_rgb(COLORS["line"]), 8, False, 32)
    node((500, 500, 1070, 700), COLORS["blue"], COLORS["line"], ["Fine–Gray 主模型", "外层患者级内部验证", "折内 MICE / 标准化 / 收缩"])
    node((1220, 500, 1790, 700), COLORS["amber"], COLORS["amber_line"], ["性能与稳健性", "区分度 · Brier · 校准 · DCA", "1 h person-period / 预设敏感性"])
    node((1920, 500, 2270, 700), COLORS["teal"], COLORS["teal_line"], ["锁定模型", "系数 + 预处理", "阈值与版本封存"])
    d.line((2005, 375, 2005, 455), fill=hex_rgb(COLORS["line"]), width=8)
    d.line((2005, 455, 785, 455), fill=hex_rgb(COLORS["line"]), width=8)
    draw_arrow(d, (785,455), (785,500), hex_rgb(COLORS["line"]), 8, False, 32)
    draw_arrow(d, (1070,600), (1220,600), hex_rgb(COLORS["line"]), 8, False, 32)
    draw_arrow(d, (1790,600), (1920,600), hex_rgb(COLORS["line"]), 8, False, 32)
    d.line((2095,700,2095,835), fill=hex_rgb(COLORS["teal_line"]), width=8)
    d.line((2095,835,410,835), fill=hex_rgb(COLORS["teal_line"]), width=8)
    draw_arrow(d, (410,835), (410,1015), hex_rgb(COLORS["teal_line"]), 8, True, 32)
    draw_centered(d, (500, 760, 2000, 830), "锁模后原样应用；再校准与原始外部验证分开报告", font(27), muted)
    node((500, 930, 2270, 1100), COLORS["grey"], COLORS["grey_line"],
         ["外部验证与适用性评估", "同口径时间窗、结局可测性、校准与临床净获益；报告双库测量差异", "early sepsis 为预设亚组；不因外部结果回改开发定义"], 32, 25)
    rounded(d, (80, 1180, 2270, 1325), hex_rgb(COLORS["white"]), hex_rgb(COLORS["grey_line"]), 28, 4)
    d.text((120, 1210), "全程控制", font=font(29, True), fill=ink)
    d.text((315, 1210), "变量字典与精确语义 allowlist · 单位/比较符号/迟到结果审计 · 患者级重抽样",
           font=font(25), fill=muted)
    d.text((120, 1260), "TRIPOD+AI 与 PROBAST+AI；最终人数、事件数和性能来自同一冻结运行。",
           font=font(25), fill=muted)
    im.save(path, dpi=(300, 300))


def set_pptx_text(shape, text, size=14, bold=False, color=COLORS["ink"], align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    for idx, line in enumerate(text.split("\n")):
        if idx:
            p = tf.add_paragraph()
            p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = "PingFang SC"
        run.font.size = PptxPt(size if idx == 0 else max(size-2, 9))
        run.font.bold = bold if idx == 0 else False
        run.font.color.rgb = PptxRGBColor(*hex_rgb(color if idx == 0 else COLORS["muted"]))


def pptx_round(slide, x, y, w, h, fill, outline, text, size=13):
    sh = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, PptxCm(x), PptxCm(y), PptxCm(w), PptxCm(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = PptxRGBColor(*hex_rgb(fill))
    sh.line.color.rgb = PptxRGBColor(*hex_rgb(outline)); sh.line.width = PptxPt(1.2)
    set_pptx_text(sh, text, size=size, bold=True)
    return sh


def pptx_arrow(slide, x1, y1, x2, y2, color=COLORS["line"], dashed=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, PptxCm(x1), PptxCm(y1), PptxCm(x2), PptxCm(y2))
    c.line.color.rgb = PptxRGBColor(*hex_rgb(color)); c.line.width = PptxPt(1.8)
    if dashed:
        c.line.dash_style = 2
    ln = c._element.spPr.ln
    head = PptxOxmlElement("a:headEnd"); head.set("type", "triangle"); ln.append(head)
    return c


def build_pptx(path: Path):
    prs = Presentation()
    prs.slide_width = PptxCm(18)
    prs.slide_height = PptxCm(10.5)
    blank = prs.slide_layouts[6]

    # Slide 1: timeline
    s = prs.slides.add_slide(blank)
    title = s.shapes.add_textbox(PptxCm(0.5), PptxCm(0.25), PptxCm(16.8), PptxCm(0.8))
    set_pptx_text(title, "Landmark 时间轴与信息边界", 18, True, align=PP_ALIGN.LEFT)
    sub = s.shapes.add_textbox(PptxCm(0.5), PptxCm(0.95), PptxCm(16.8), PptxCm(0.5))
    set_pptx_text(sub, "用前 12 h 可得信息预测随后 48 h 风险；活着离开 index ICU 为竞争事件", 10, False, COLORS["muted"], PP_ALIGN.LEFT)
    pptx_round(s, 0.7, 1.9, 7.9, 1.05, COLORS["teal"], COLORS["teal_line"], "DHF 表型支持窗 [T0-24 h, T12)\nHF 锚点 + 失代偿证据 + 客观/管理支持", 12)
    pptx_round(s, 2.5, 3.25, 6.1, 0.85, COLORS["blue"], COLORS["line"], "预测器窗 [T0, T12)", 12)
    pptx_round(s, 8.6, 3.25, 6.7, 0.85, COLORS["amber"], COLORS["amber_line"], "风险窗 [T12, min(T60, 活着出 ICU))", 11)
    pptx_arrow(s, 2.2, 4.45, 15.9, 4.45)
    for x in (2.5, 8.6, 15.3):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, PptxCm(x), PptxCm(1.55), PptxCm(x), PptxCm(7.25))
        c.line.color.rgb = PptxRGBColor(*hex_rgb(COLORS["ink"])); c.line.width = PptxPt(1.2)
    labels = [(1.4,"T0\nindex ICU 实际入科"),(7.5,"T12\n仍存活且留在 index ICU"),(14.2,"T60\n行政随访上限")]
    for x, txt in labels:
        box=s.shapes.add_textbox(PptxCm(x),PptxCm(7.25),PptxCm(2.3),PptxCm(1.0)); set_pptx_text(box,txt,11,True)
    pptx_round(s, 9.0, 5.15, 3.5, 1.0, COLORS["rose"], COLORS["rose_line"], "目标事件\n血流动力学恶化或 ICU 内死亡", 10)
    pptx_round(s, 12.8, 5.15, 2.8, 1.0, COLORS["teal"], COLORS["teal_line"], "竞争事件\n活着离开 index ICU", 10)
    note=s.shapes.add_textbox(PptxCm(0.6),PptxCm(9.45),PptxCm(16.8),PptxCm(0.45)); set_pptx_text(note,"注：后写文书可用于重建时间，但不得把 T12 后可见信息输入预测器。",9,False,COLORS["muted"],PP_ALIGN.LEFT)

    # Slide 2: pipeline
    s = prs.slides.add_slide(blank)
    title = s.shapes.add_textbox(PptxCm(0.5), PptxCm(0.2), PptxCm(17), PptxCm(0.7)); set_pptx_text(title,"双库预测模型技术路线",18,True,align=PP_ALIGN.LEFT)
    sub = s.shapes.add_textbox(PptxCm(0.5), PptxCm(0.85), PptxCm(17), PptxCm(0.45)); set_pptx_text(sub,"先冻结对象、时间和结局，再开发模型并进行锁模后外部验证",10,False,COLORS["muted"],PP_ALIGN.LEFT)
    nodes=[
        (0.5,1.55,2.4,1.3,COLORS["blue"],COLORS["line"],"MIMIC-IV\n开发 + 内部验证"),
        (3.5,1.55,2.4,1.3,COLORS["grey"],COLORS["grey_line"],"键与 ICU 时间轴\n唯一 episode"),
        (6.5,1.55,2.6,1.3,COLORS["teal"],COLORS["teal_line"],"DHF 多域表型\nA + B + C"),
        (9.7,1.55,2.6,1.3,COLORS["amber"],COLORS["amber_line"],"T12 风险集\n存活并留在 ICU"),
        (12.9,1.55,4.2,1.3,COLORS["rose"],COLORS["rose_line"],"预测器 + 三态结局\n目标 / 活着出 ICU / 删失"),
    ]
    for n in nodes: pptx_round(s,*n,size=10)
    for x1,x2 in [(2.9,3.5),(5.9,6.5),(9.1,9.7),(12.3,12.9)]: pptx_arrow(s,x1,2.2,x2,2.2)
    pptx_round(s, 3.5, 4.0, 4.2, 1.45, COLORS["blue"], COLORS["line"], "Fine–Gray 主模型\n患者级内部验证；折内 MICE / 标准化", 11)
    pptx_round(s, 8.6, 4.0, 4.2, 1.45, COLORS["amber"], COLORS["amber_line"], "性能与稳健性\n区分度 · Brier · 校准 · DCA", 11)
    pptx_round(s, 14.0, 4.0, 3.1, 1.45, COLORS["teal"], COLORS["teal_line"], "锁定模型\n系数 + 预处理", 11)
    pptx_arrow(s, 15.0,2.85,5.6,4.0); pptx_arrow(s,7.7,4.73,8.6,4.73); pptx_arrow(s,12.8,4.73,14.0,4.73)
    pptx_round(s, 0.5, 7.0, 2.4, 1.25, COLORS["teal"], COLORS["teal_line"], "本院数据库\n锁模后外部验证", 10)
    pptx_round(s, 3.5, 7.0, 13.6, 1.25, COLORS["grey"], COLORS["grey_line"], "外部验证与适用性评估\n同口径时间窗与结局可测性；再校准与原始外部验证分开报告；early sepsis 预设亚组", 10)
    pptx_arrow(s,15.5,5.45,2.9,7.6,COLORS["teal_line"],True); pptx_arrow(s,2.9,7.62,3.5,7.62,COLORS["teal_line"],True)
    foot=s.shapes.add_textbox(PptxCm(0.6),PptxCm(9.1),PptxCm(16.8),PptxCm(0.85)); set_pptx_text(foot,"全程控制：变量字典、精确语义 allowlist、单位/比较符号/迟到结果审计、患者级重抽样、TRIPOD+AI 与 PROBAST+AI。最终数字来自同一冻结运行。",9,False,COLORS["muted"],PP_ALIGN.LEFT)
    prs.save(path)


def write_figure_contracts():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    figures = {
        "figure1_timeline": {
            "question": "如何用明确的时间边界避免预测器泄漏，并正确处理 ICU 出科竞争事件？",
            "claim": "T12 是预测时点；预测器仅来自 [T0,T12)，风险随访止于 T60 或活着出 index ICU。",
            "source": ["project_control/STATISTICAL_ANALYSIS_PLAN_DHF_PREDICTION_V1.md §2",
                       "project_control/DHF_PROJECT_REPRODUCIBILITY_CONTRACT_V1.md §2-3"],
            "evidence_status": "schematic",
            "prohibited_implication": "不得把 T12 后资料作为预测器，亦不得把活着出 ICU 当作普通无事件。",
            "reading_path": ["T0", "T12", "T60"],
        },
        "figure2_pipeline": {
            "question": "如何从双库原始资料形成可锁定、可内部验证并可外部验证的预测模型？",
            "claim": "必须先冻结时间轴、DHF 表型、风险集和三态结局，再开发 Fine–Gray 模型并执行锁模后外部验证。",
            "source": ["project_control/RESEARCH_LOGIC_CHAIN_20260916.md",
                       "project_control/PROPOSAL_METHODS_DRAFT_20260915.md",
                       "project_control/STATISTICAL_ANALYSIS_PLAN_DHF_PREDICTION_V1.md"],
            "evidence_status": "schematic/prospective",
            "prohibited_implication": "流程图不表示最终队列、结局或模型已冻结。",
            "reading_path": ["MIMIC", "phenotype", "riskset", "outcome", "model", "locked external validation"],
        },
    }
    for name, spec in figures.items():
        d = ASSET_DIR / name
        d.mkdir(parents=True, exist_ok=True)
        (d / "content_ledger.json").write_text(json.dumps({"schema_version":"1.2","figure_id":name,"claims":[spec]}, ensure_ascii=False, indent=2), encoding="utf-8")
        (d / "layout_blueprint.json").write_text(json.dumps({
            "schema_version":"1.2","figure_id":name,"blueprint_revision":1,"mode":"flowchart",
            "question":spec["question"],"narrative_job":spec["claim"],
            "decision_unlocked":"允许读者判断时间边界、模型估计对象和双库角色是否一致",
            "visual_anchor":{"panel_id":"main","reason":"单页主流程是唯一科学叙事"},
            "reading_path":spec["reading_path"],"groups":[{"id":"main","panels":["main"],"role":"schematic","emphasis":"primary"}],
            "panels":[{"id":"main","group_id":"main","evidence_role":"schematic","priority":"anchor","preferred_aspect":"wide","shared_scale_group":None,"legend_owner":"main","source_data_required":False}],
            "shared_scale_groups":[],"layout_rationale":"横向阅读；时间或分析阶段决定面积，不按装饰性等分。"
        }, ensure_ascii=False, indent=2), encoding="utf-8")
        (d / "flowchart_spec.json").write_text(json.dumps({
            "schema_version":"1.2","figure_id":name,"stages":spec["reading_path"],
            "edge_semantics":{"solid":"预设主流程或时间推进","dashed":"竞争事件、拟执行或锁模后分支"},
            "evidence_status":spec["evidence_status"],"source_anchors":spec["source"]
        }, ensure_ascii=False, indent=2), encoding="utf-8")
        (d / "source_data_manifest.json").write_text(json.dumps({
            "schema_version":"1.2","figure_id":name,"source_data_required":False,
            "reason":"纯研究设计示意图，不包含患者级数据、测量图像、模型性能或定量结果。",
            "delivery_outputs":[f"{name}.png",f"{name}.svg"]
        }, ensure_ascii=False, indent=2), encoding="utf-8")
        (d / "qa_note.md").write_text(
            "# QA note\n\n- 内容：与研究合同逐项核对，所有元素均为 schematic/prospective。\n"
            "- 版式：180 mm 宽，中文正文在打印尺寸下可读；颜色同时以边框、线型和直接标签冗余编码。\n"
            "- 模型概念生图：`not_run`，当前运行环境未暴露可核验的图像生成端点；直接依据已检查的 Nature 风格矢量模板重建。\n"
            "- 输出：PNG 用于 DOCX；SVG 与原生 PPTX 用于后续编辑。\n",
            encoding="utf-8")


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill.replace("#", ""))


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m, v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}"); tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B8C4CF", size="6"):
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)
        borders.append(node)
    tbl_pr.append(borders)


def apply_run_style(run, size=12, bold=False, italic=False, color=None, underline=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*hex_rgb(color))
    if underline is not None:
        run.font.underline = underline


def format_paragraph(p, *, first_line=True, line=1.5, after=4, before=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep=False):
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.first_line_indent = Cm(0.74) if first_line else None
    pf.keep_with_next = keep


def add_paragraph(doc, text, size=12, bold=False, first_line=True, line=1.5, after=4, before=0,
                  align=WD_ALIGN_PARAGRAPH.JUSTIFY, style="Normal", keep=False, color=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    apply_run_style(r, size=size, bold=bold, color=color)
    format_paragraph(p, first_line=first_line, line=line, after=after, before=before, align=align, keep=keep)
    return p


def add_main_heading(doc, text):
    return add_paragraph(doc, text, size=12, bold=True, first_line=False, before=6, after=5,
                         align=WD_ALIGN_PARAGRAPH.JUSTIFY, style="HTML 预设格式1", keep=True)


def add_subheading(doc, text):
    return add_paragraph(doc, text, size=12, bold=True, first_line=False, before=4, after=4,
                         align=WD_ALIGN_PARAGRAPH.LEFT, keep=True)


def add_numbered(doc, label, text):
    p = doc.add_paragraph(style="Normal")
    r = p.add_run(label)
    apply_run_style(r, 12, True)
    r = p.add_run(text)
    apply_run_style(r, 12, False)
    format_paragraph(p, first_line=False, line=1.5, after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    return p


def add_caption(doc, text):
    return add_paragraph(doc, text, size=10.5, bold=False, first_line=False, line=1.2, after=5,
                         align=WD_ALIGN_PARAGRAPH.CENTER, color=COLORS["muted"])


def add_picture(doc, path, width_cm, alt_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    for docpr in p._p.xpath(".//wp:docPr"):
        docpr.set("descr", alt_text)
    return p


def add_progress_table(doc):
    rows = [
        ("院内候选与时间门控", "成人 index ICU 候选分母 8,385；完成成人、唯一键、T12 偏移、出 ICU 顺序和可用时间窗 QC", "候选分母，不是最终 DHF 队列"),
        ("院内证据整合", "284,049 条去重证据、20,025 条目标医嘱；3,720 人有同窗 BNP，6,432/8,385 有同窗可解析乳酸", "覆盖率与工作量审计，不是阳性率或结局结果"),
        ("院内病例复核", "形成 235 例优先队列、91 例 T12 观察链、42 例 HF 锚点队列；36 份空模板剔除、1 份矛盾文书隔离", "工作清单及 AI 预审仍待临床裁决"),
        ("MIMIC 影像", "主窗口 7,828/7,828 份报告完成导出与对账，涉及 3,886 个有报告 stay；建立 300 条 round 1 和 60 条盲法 round 2 标注包", "300 条尚待临床人工最终确认"),
        ("实验室数据合同", "PostgreSQL 全库 37,778,198 条 classified、37,732,919 条 eligible；45,279 条错误体液/反向时间/未知单位记录隔离，12 个硬门通过", "数据清洗 QC，不代表正式队列、事件或模型结果"),
        ("跨平台复核", "固定 5,549-stay 审计快照完成 BigQuery 聚合复核；14 个概念、9 个硬门通过，91/91 项回归与静态门通过", "审计 SQL 仍为 AUDIT_ONLY，allow_final_run=false"),
        ("模型状态", "旧模型与历史宽口径运行仅证明代码和流程可执行", "正式 Fine–Gray、内部验证与本院外部验证均未运行/未冻结"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Normal Table"
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(3.0), Cm(8.6), Cm(4.0)]
    hdr = table.rows[0].cells
    for i, txt in enumerate(["工作模块", "截至 2026-09-20 的已验证进展", "解释边界"]):
        hdr[i].text = txt
        hdr[i].width = widths[i]
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(hdr[i], COLORS["blue_strong"])
        set_cell_margins(hdr[i])
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: apply_run_style(r, 9.5, True)
    for row in rows:
        cells = table.add_row().cells
        for i, txt in enumerate(row):
            cells[i].text = txt
            cells[i].width = widths[i]
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            if i == 0: set_cell_shading(cells[i], COLORS["grey"])
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs: apply_run_style(r, 9.2, bold=(i == 0))
    add_caption(doc, "表1  当前研究进展、可核验证据及其解释边界")


def remove_old_body(doc):
    body = doc._element.body
    keep = {id(p._p) for p in doc.paragraphs[:10]}
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        if id(child) not in keep:
            body.remove(child)


def prune_unused_image_relationships(doc):
    used = set()
    embed_attr = qn("r:embed")
    for element in doc._element.iter():
        rid = element.get(embed_attr)
        if rid:
            used.add(rid)
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype == DOCX_RT.IMAGE and rid not in used:
            doc.part.drop_rel(rid)


def update_cover(doc):
    # Preserve the school cover artwork, paragraph formatting and section break.
    doc.paragraphs[6].runs[1].text = "2026 年 _"
    doc.paragraphs[6].runs[2].text = "9"
    doc.paragraphs[6].runs[3].text = "_ 月 _"
    doc.paragraphs[6].runs[4].text = "20"
    doc.paragraphs[6].runs[5].text = "_ 日"
    p = doc.paragraphs[7]
    p.runs[1].text = COVER_TITLE
    p.runs[1].font.size = Pt(10.5)
    p.runs[1]._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    p.paragraph_format.line_spacing = 1.15


def add_body_content(doc):
    add_main_heading(doc, "（一）选题依据")
    add_main_heading(doc, "1. 课题来源：导师研究课题的一部分")
    add_main_heading(doc, "2. 课题的研究意义、国内外研究现状分析")
    add_subheading(doc, "2.1 课题的研究意义")
    add_paragraph(doc, "失代偿性心力衰竭（decompensated heart failure，DHF）患者进入重症监护病房后，循环状态、容量负荷与器官灌注可在短时间内快速变化。临床团队需要在完成初始评估和治疗后，识别未来数十小时内可能需要循环支持升级或发生死亡的高风险患者。既有心源性休克分期和心衰相关休克共识强调早期识别与重复评估，但休克是动态临床状态，单一诊断编码、血压阈值、利钠肽或一次心超均不足以重建其发生过程。[1,2,6,7]")
    add_paragraph(doc, "本研究将预测时点设置为 index ICU 实际入科后 12 h（T12）：一方面，前 12 h 已累积生命体征、实验室、治疗和支持信息；另一方面，仍保留随后 48 h 的临床干预窗口。主要预测对象不是狭义或经人工确诊的心源性休克，而是预设的治疗升级相关 ICU 血流动力学恶化或 ICU 内死亡；同时把活着离开 index ICU 作为竞争事件。该设定更贴近 ICU 风险管理，也可避免把已离开 ICU 的患者简单视为与持续留在 ICU 者具有相同观察机会。")
    add_paragraph(doc, "研究采用“公共数据库开发与内部验证—锁定模型—单中心外部验证”的路径。MIMIC-IV 提供可复现的重症电子病历数据，本院资料可检验模型在不同检查选择、时间记录和临床流程下的可迁移性。[8,9] 若研究顺利完成，可形成一套边界清晰、可审计并能够进一步转化为院内早期风险提示的预测框架。")
    add_subheading(doc, "2.2 国内外研究现状")
    add_paragraph(doc, "既往研究已经证明，利用早期电子病历信息预测心源性休克或心衰恶化具有可行性。Hu 等在心脏 ICU 混合人群中使用经医师裁决的 onset 时间开发动态休克风险评分；Beer 等前瞻性分析急性心衰患者从病情加重至休克的危险因素；Rahman 等则探索了急性失代偿性心衰住院期间的实时风险监测。[3–5] 这些工作提示，时间更新的生理、实验室和治疗信号能够在临床明确恶化前提供风险信息。")
    add_paragraph(doc, "然而，现有证据仍存在三方面不足。第一，研究对象常为急性冠脉综合征、心脏 ICU 混合人群或宽口径心衰住院患者，DHF 操作性表型、index ICU episode 与预测时点之间的关系未必清楚。第二，不少研究采用死亡、广义 worsening heart failure 或固定二分类结局，未充分处理活着出 ICU 对后续 ICU 事件观察的竞争作用。第三，文本否定、不确定陈述、检查选择性、结果报告延迟、医嘱与真实执行差异等常规医疗数据误差，可能导致错纳、时间泄漏和外部验证口径不一致。[9–14]")
    add_paragraph(doc, "本课题据此将科学问题拆解为连续六步：候选分母、ICU 时间轴、DHF 多域表型、T12 风险集、三态结局和预测验证。每一步只使用前一步输出，最终性能不能反向修改人群、时间窗、变量或阈值。文本处理借鉴 NegEx 的局部否定和 CheXpert 的不确定标签思想，但本研究的中文规则与 AI 预审核均须以独立临床参照校准，不能自称已验证诊断工具。[10,11]")

    add_main_heading(doc, "3. 主要参考文献")
    references = [
        "[1] Bozkurt B, Coats AJS, Tsutsui H, et al. Universal definition and classification of heart failure. Eur J Heart Fail. 2021;23:352-380. doi:10.1002/ejhf.2115.",
        "[2] McDonagh TA, Metra M, Adamo M, et al. 2021 ESC Guidelines for the diagnosis and treatment of acute and chronic heart failure. Eur Heart J. 2021;42:3599-3726. doi:10.1093/eurheartj/ehab368.",
        "[3] Hu Y, Lui A, Goldstein M, et al. Development and external validation of a dynamic risk score for early prediction of cardiogenic shock in cardiac intensive care units using machine learning. Eur Heart J Acute Cardiovasc Care. 2024;13:472-480. doi:10.1093/ehjacc/zuae037.",
        "[4] Beer BN, Kellner C, Sundermeyer J, et al. Prediction of cardiac worsening through to cardiogenic shock in patients with acute heart failure. ESC Heart Fail. 2024;11:2249-2258. doi:10.1002/ehf2.14792.",
        "[5] Rahman F, Finkelstein N, Alyakin A, et al. Using machine learning for early prediction of cardiogenic shock in patients with acute heart failure. J Soc Cardiovasc Angiogr Interv. 2022;1:100308. doi:10.1016/j.jscai.2022.100308.",
        "[6] Naidu SS, Baran DA, Jentzer JC, et al. SCAI SHOCK Stage Classification Expert Consensus Update. J Am Coll Cardiol. 2022;79:933-946. doi:10.1016/j.jacc.2022.01.018.",
        "[7] Kanwar MK, Billia F, Randhawa V, et al. Heart failure related cardiogenic shock: an ISHLT consensus conference content summary. J Heart Lung Transplant. 2024;43:189-203. doi:10.1016/j.healun.2023.09.014.",
        "[8] Johnson AEW, Bulgarelli L, Shen L, et al. MIMIC-IV, a freely accessible electronic health record dataset. Sci Data. 2023;10:1. doi:10.1038/s41597-022-01899-x.",
        "[9] Benchimol EI, Smeeth L, Guttmann A, et al. The REporting of studies Conducted using Observational Routinely-collected health Data (RECORD) statement. PLoS Med. 2015;12:e1001885. doi:10.1371/journal.pmed.1001885.",
        "[10] Chapman WW, Bridewell W, Hanbury P, Cooper GF, Buchanan BG. A simple algorithm for identifying negated findings and diseases in discharge summaries. J Biomed Inform. 2001;34:301-310. doi:10.1006/jbin.2001.1029.",
        "[11] Irvin J, Rajpurkar P, Ko M, et al. CheXpert: a large chest radiograph dataset with uncertainty labels and expert comparison. Proc AAAI Conf Artif Intell. 2019;33:590-597. doi:10.1609/aaai.v33i01.3301590.",
        "[12] Fine JP, Gray RJ. A proportional hazards model for the subdistribution of a competing risk. J Am Stat Assoc. 1999;94:496-509. doi:10.1080/01621459.1999.10474144.",
        "[13] Collins GS, Dhiman P, Andaur Navarro CL, et al. TRIPOD+AI statement: updated guidance for reporting clinical prediction models that use regression or machine learning methods. BMJ. 2024;385:e078378. doi:10.1136/bmj-2023-078378.",
        "[14] Moons KGM, Damen JAA, Kaul T, et al. PROBAST+AI: an updated quality, risk of bias, and applicability assessment tool for prediction models using regression or artificial intelligence methods. BMJ. 2025. doi:10.1136/bmj-2024-082505.",
        "[15] Riley RD, Ensor J, Snell KIE, et al. Calculating the sample size required for developing a clinical prediction model. BMJ. 2020;368:m441. doi:10.1136/bmj.m441.",
        "[16] White IR, Royston P, Wood AM. Multiple imputation using chained equations: issues and guidance for practice. Stat Med. 2011;30:377-399. doi:10.1002/sim.4067.",
        "[17] Vickers AJ, Elkin EB. Decision curve analysis: a novel method for evaluating prediction models. Med Decis Making. 2006;26:565-574. doi:10.1177/0272989X06295361.",
        "[18] Steyerberg EW. Clinical Prediction Models. 2nd ed. Cham: Springer; 2019.",
    ]
    for ref in references:
        add_paragraph(doc, ref, size=10.5, first_line=False, line=1.2, after=1.5, align=WD_ALIGN_PARAGRAPH.LEFT)

    add_main_heading(doc, "（二）研究方案")
    add_main_heading(doc, "1. 研究目标、内容和拟解决的关键问题")
    add_subheading(doc, "1.1 研究目标")
    add_paragraph(doc, "总体目标是在成人 index ICU 的 DHF 操作性表型患者中，利用 T12 前可获得的信息，建立并验证未来 48 h 治疗升级相关 ICU 血流动力学恶化或 ICU 内死亡的个体风险预测模型。研究属于预后预测，不估计利尿剂、血管活性药、机械支持或其他治疗的因果效应。")
    add_numbered(doc, "① ", "建立可追溯的 DHF 多域操作性表型、index ICU 时间轴和 T12 风险集，明确每例患者的证据来源、时间和不确定性。")
    add_numbered(doc, "② ", "在 MIMIC 中开发以 Fine–Gray 为主的紧凑预测模型并进行严格内部验证，评估区分度、总体误差、校准和临床净获益。")
    add_numbered(doc, "③ ", "锁定预处理、变量、系数和风险计算后，在本院队列开展原样外部验证；如需再校准，与原始外部验证结果分开报告。")
    add_numbered(doc, "④ ", "通过 1 h person-period、表型/时间/缺失/结局敏感性分析及 early sepsis 预设亚组评价结果稳健性与适用范围。")

    add_subheading(doc, "1.2 研究内容")
    add_numbered(doc, "① 队列与时间轴：", "从成人候选分母重建唯一患者、住院与 ICU episode，确定 T0、T12、T60、再入 ICU、出 ICU 和死亡的先后关系。")
    add_numbered(doc, "② DHF 表型：", "按 A（本次 HF 锚点）、B（独立失代偿/充血证据）、C（心超、利钠肽或患者特异性管理支持）记录证据及替代解释；C 域不能单独确诊。")
    add_numbered(doc, "③ 风险集与结局：", "纳入 T12 仍存活并留在 index ICU、且未处于预设既有严重状态代理的患者；重建目标事件、活着出 ICU 竞争事件和行政删失。")
    add_numbered(doc, "④ 预测器：", "仅使用 [T0,T12) 内已可获得的人口学、入科特征、生命体征、灌注/肾功能、呼吸支持、液体与尿量及支持治疗信息。")
    add_numbered(doc, "⑤ 模型与验证：", "在 MIMIC 完成开发和内部验证，在本院完成锁模后外部验证，并按照 TRIPOD+AI 与 PROBAST+AI 报告透明度、偏倚及适用性。[13,14]")

    add_subheading(doc, "1.3 拟解决的关键问题")
    add_numbered(doc, "① ", "如何避免把慢性稳定心衰、单项 BNP 升高、孤立心超异常、非心源性低氧或模板措辞误纳为本次 DHF。")
    add_numbered(doc, "② ", "如何把同一次住院的多次 ICU episode 正确切分，并确保报告、检验和文书在预测时点前真实可用，降低时间泄漏。")
    add_numbered(doc, "③ ", "如何在存在活着出 ICU 竞争事件的情况下定义和评价 48 h 累积发生风险，而不把不同观察机会混为固定二分类问题。")
    add_numbered(doc, "④ ", "如何处理两库检查选择、结果时间、单位、缺失和医嘱执行代理差异，使外部验证评价的是同一临床问题。")
    add_numbered(doc, "⑤ ", "如何在最终事件数尚未冻结的情况下控制模型复杂度，避免以旧宽口径队列或单因素筛选决定最终变量。")

    add_main_heading(doc, "2. 拟采取的研究方法、技术路线、实施方案及可行性分析")
    add_subheading(doc, "2.1 研究设计与数据来源")
    add_paragraph(doc, "本研究为回顾性双数据库预后预测模型研究。MIMIC-IV 作为开发与内部验证数据库；本院数据库作为模型锁定后的单中心外部验证数据库。每名患者原则上仅以首次满足预设条件的 index ICU episode 进入分析；同次住院多次 ICU 段分别重建，禁止把后续 ICU 段资料拼入前一段。MIMIC 使用多域 DHF 操作性表型，心超作为支持和敏感性层；本院现有候选分母源自成人及床旁心超选择，需明确报告其选择边界。")
    add_paragraph(doc, "研究将遵守 MIMIC 数据使用协议及院内数据安全要求。患者级原始资料不离开受控环境，公开工件仅包括去标识化代码、数据字典、聚合审计和必要的合成示例。院内伦理批件号、覆盖年月及知情同意豁免状态在取得正式文件后据实补录，不在本报告中预设。")

    add_subheading(doc, "2.2 时间锚点、纳排标准与信息边界")
    add_paragraph(doc, "T0 为 index ICU 实际入科时间，优先使用护理到达记录或正文中明确的 ICU 转入事件；普通住院时间、文书创建时间和“拟转出”不自动等同实际 ICU 事件。DHF 表型支持窗为 [T0−24 h,T12)，预测器窗为 [T0,T12)，风险窗从 T12 起至 min(T60, 活着离开 index ICU)。T60 指 T0 后 60 h，因此 T12 后风险时长最多 48 h。")
    add_paragraph(doc, "纳入条件为：成人、唯一 index ICU episode、满足冻结版 DHF 操作性表型、T12 仍存活且留在 index ICU、预测器和结局可按合同评估。排除条件包括：键或时间轴无法解析、T12 前已死亡或离开 index ICU、处于预设既有严重状态代理、关键观察完全不可判定，或其他冻结版合同规定的排除情况。缺乳酸、缺检查或缺文书不自动证明阴性，而应记录相应缺失和观察完整性状态。")
    add_picture(doc, ASSET_DIR / "figure1_timeline" / "figure1_timeline.png", 15.6,
                "T0-T12-T60 landmark 时间轴，显示表型、预测器和竞争风险结局窗口")
    add_caption(doc, "图1  Landmark 时间轴与信息边界。T12 为预测时点；目标事件、活着出 ICU 与行政删失按互斥状态构建。")

    add_subheading(doc, "2.3 DHF 操作性表型与文本审核")
    add_paragraph(doc, "DHF 表型由三类证据构成。A 域要求本次 episode 相关的 HF 诊断或与失代偿相联系的结构/功能异常；B 域要求独立的失代偿或充血表现，如呼吸困难、湿啰音、水肿、客观肺充血或与心源性机制一致的低灌注；C 域记录心超、BNP/NT-proBNP 及患者特异性管理支持。A 与 B 是主干，C 仅作支持。肺炎、ARDS、感染、出血、术后和肾功能不全等替代解释逐例记录是否足以解释当前表现。[1,2]")
    add_paragraph(doc, "文本规则识别否定、不确定、既往史、风险告知和模板语句，并保留来源、行号、时间和关键原文。AI 仅用于原文预审核和工作队列排序；最终临床参照、阳性预测值、敏感度及评阅者一致性必须依赖独立临床审核与正确抽样。第二位标注者的盲法复核方可用于评阅者间一致性，单一标注者重复复核不能代替独立 kappa。")

    add_subheading(doc, "2.4 主要结局、竞争事件与可观察性")
    add_paragraph(doc, "主要结局为 T12 后至风险窗结束前发生的预设治疗升级相关 ICU 血流动力学恶化或 ICU 内死亡。治疗升级要求具有明确的 T12 后时间戳，并相对 T12 基线出现符合冻结合同的循环支持强化；具体组成、强度阈值和持续时间须在双库可测性核对后预先冻结。死亡必须定位于 index ICU 段。活着离开 index ICU 为竞争事件；T60 仍在 ICU 且未发生目标事件者为行政删失；观察不完整或关键支持强度无法判定者不得默认为无事件。")
    add_paragraph(doc, "本院在 eMAR 暂不可得时，以未作废医嘱的开停时间和给药途径作为治疗暴露代理，文书与护理记录用于交叉核对。医嘱不等于实际给药，药品总量不能推导泵速或去甲肾上腺素等效剂量。取得执行级记录后，将在同一冻结队列中进行预设敏感性重跑。")

    add_subheading(doc, "2.5 预测器、数据清洗与缺失处理")
    add_paragraph(doc, "候选预测器按临床合理性、T12 前可得性、两库可迁移性、测量质量和模型复杂度预先确定，不使用单因素 P 值筛选。主要域包括人口学与入科方式、生命体征、酸碱与灌注、肾功能、电解质、血细胞、呼吸支持、尿量/液体平衡和前 12 h 支持治疗结构。与结局定义重叠、T12 后可得或从结局反推的变量进入黑名单。")
    add_paragraph(doc, "清洗保留原始字符串、原始单位、采样/报告/存储时间、比较符号、异常标志、来源表和行号。结果可用时间取真实可见时间；采样在窗内但 T12 后才出结果的记录不进入预测器。`>x`、`<x` 和区间结果保留上下界及删失标志，不当作精确值、0 或中点；只有整个区间位于阈值一侧时才作确定性阈值判断。单位按项目字典和精确语义 allowlist 处理，错误体液、单位不符和无法解释值进入隔离审计，不静默删除或补零。")
    add_paragraph(doc, "缺失分为结构性、测量性和信息性。只有字典明确“无记录即未暴露”的结构性缺失可填 0；常规连续变量采用外层训练折内多重插补，计划 m=20，并以 predictive mean matching 为主。[16] 插补、编码、标准化、非线性转换和变量选择全部在训练折内拟合。心超和影像属于选择性检查，保留 not_done/unavailable 状态，不对未检查者虚构数值。")

    add_subheading(doc, "2.6 模型开发、内部验证与外部验证")
    add_paragraph(doc, "主模型采用 Fine–Gray 亚分布风险模型，直接估计存在活着出 ICU 竞争事件时的 48 h 目标事件累积发生风险。[12] 最终队列和事件数冻结后，按 Riley 框架重新评估可支持的有效参数、收缩程度与校准精度，不机械使用固定每变量事件数规则。[15] 若事件数不足，将按预设临床域压缩参数或加强收缩，而不是在多种算法中挑选最好看的结果。")
    add_paragraph(doc, "内部验证以患者为重抽样单位。计划采用外层 5 折生成 out-of-fold 风险；若纳入惩罚或超参数选择，则仅在外层训练折内进行内层交叉验证。报告 48 h 竞争风险区分度、Brier 分数、校准截距/斜率与校准图、风险分层及决策曲线，并给出 95% 置信区间。[17,18] 1 h person-period 模型用于检验时间离散化结构，cause-specific 分析和固定时点二分类仅作为明确不同估计目标的补充，不能冒充主 Fine–Gray 结果。")
    add_paragraph(doc, "最终模型在全部开发资料上按锁定规则重拟合，并封存变量映射、插补器、转换参数、系数、版本和风险计算。外部验证在本院原样应用锁定模型，报告区分度、总体误差、校准与临床净获益；若进行截距或斜率再校准，应与未经更新的外部验证结果分开呈现。任何外部结果不得反向修改开发队列、预测器或结局定义。")

    add_subheading(doc, "2.7 敏感性分析与偏倚控制")
    add_paragraph(doc, "预设敏感性分析包括：DHF 多域主层与影像/心超严格层；报告时间与正式采样/执行时间；比较符号与删失值处理；折内 MICE、缺失指示和完全病例诊断；医嘱代理与执行级给药；目标事件组成与持续时间；前 12 h 既有严重状态代理；同窗心超选择；以及 early sepsis 亚组。研究全程以 TRIPOD+AI 规范报告，以 PROBAST+AI 审查参与者选择、预测器、结局和分析偏倚。[13,14]")

    add_subheading(doc, "2.8 技术路线")
    add_picture(doc, ASSET_DIR / "figure2_pipeline" / "figure2_pipeline.png", 15.6,
                "双库DHF表型、风险集、结局、Fine-Gray内部验证与锁模外部验证技术路线")
    add_caption(doc, "图2  双库预测模型技术路线。实线为预设主流程，虚线表示锁模后外部验证路径；流程图不代表最终队列或模型已经冻结。")

    add_subheading(doc, "2.9 研究实施方案")
    add_numbered(doc, "阶段一：", "完成双库 index ICU 时间轴、DHF 多域表型的临床裁决与冻结，形成逐例纳排流转表、证据来源表和规则阴性抽样审计。")
    add_numbered(doc, "阶段二：", "在同一冻结运行中构建 T12 风险集与三态结局，核对治疗升级组成、出 ICU、死亡、删失和观察完整性。")
    add_numbered(doc, "阶段三：", "冻结紧凑候选变量与数据处理合同，在 MIMIC 进行 Fine–Gray 开发、患者级内部验证和预设敏感性分析。")
    add_numbered(doc, "阶段四：", "封存最终模型与预处理，在本院执行原样外部验证、适用性分析及必要的独立再校准。")
    add_numbered(doc, "阶段五：", "形成纳排流程图、Table 1、结局组成表、模型系数与性能表、校准/决策曲线、敏感性结果及 TRIPOD+AI/PROBAST+AI 材料。")

    add_subheading(doc, "2.10 可行性分析")
    add_paragraph(doc, "本课题已具备 MIMIC 受控数据访问、院内候选资料、SQL/Python/R 分析环境和双库变量映射基础。研究已经形成表型合同、统计分析计划、变量与语义字典、医嘱代理规范、运行登记和 fail-closed 质量门，能够在任何正式定义未冻结时阻止最终模型运行。现有工作证明数据链接、时间门控、实验室合同和审计代码可执行，但不将工程可行性夸大为临床有效性。")
    add_progress_table(doc)
    add_paragraph(doc, "综上，数据来源、分析环境与质量控制路径能够支持后续研究；当前主要阻塞不在算力，而在双库临床裁决、结局执行级可测性与正式冻结。研究将优先解决这些决定模型估计对象的问题，再进入建模，避免以算法性能掩盖标签和时间边界不确定性。")

    add_main_heading(doc, "3. 研究创新点")
    add_numbered(doc, "① 科学问题与估计对象一致：", "以 T12 landmark 后 48 h 累积发生风险为目标，把活着出 ICU 作为竞争事件，而非简单二分类。")
    add_numbered(doc, "② 多域 DHF 表型与风险集分离：", "先确认本次 DHF，再判断 T12 风险资格，避免把心衰历史、单项检查或已存在严重状态混入“随后恶化”。")
    add_numbered(doc, "③ 强化结果可用时间和语义数据合同：", "同时审计 specimen、单位、比较符号、迟到结果、episode 粒度与医嘱代理，系统控制常规医疗数据中的隐性泄漏。")
    add_numbered(doc, "④ 双库角色预先锁定：", "MIMIC 只承担开发与内部验证，本院只在锁模后外部验证；两库测量差异作为适用性边界，而非通过结果导向修改口径。")
    add_numbered(doc, "⑤ AI 使用边界可审计：", "AI 用于证据预审核和工作排序，保留原文锚点与 provenance，不把 AI 自我一致性写成临床金标准或独立阅片性能。")

    add_main_heading(doc, "4. 研究计划及预计进展")
    add_numbered(doc, "前期基础（截至 2026.09.20）：", "完成研究问题重构、双库角色与 T0/T12/T60 时间轴确定、DHF 操作性表型和统计方案、院内候选时间门控、MIMIC 主窗口影像导出、实验室数据合同及 PostgreSQL/BigQuery 聚合 QC。正式队列、三态结局和模型尚未冻结。")
    add_numbered(doc, "第一阶段（2026.09–2026.10）：", "完成 MIMIC 300 条临床确认及必要的 60 条第二标注者盲法复核；完成院内优先病例、T12 观察链、HF 锚点与规则阴性抽样裁决；冻结双库 DHF 表型。")
    add_numbered(doc, "第二阶段（2026.10–2026.11）：", "冻结 T12 风险集、三态结局与紧凑候选变量；完成事件数和模型复杂度评估，运行 Fine–Gray 开发与严格内部验证。")
    add_numbered(doc, "第三阶段（2026.11–2026.12）：", "完成 1 h person-period、缺失/时间/表型/结局敏感性分析；封存模型并在本院进行原样外部验证和适用性评估。")
    add_numbered(doc, "第四阶段（2027.01–2027.03）：", "整理学位论文和图表，完成 TRIPOD+AI 报告、PROBAST+AI 偏倚审查、代码与复现材料归档，并根据导师意见修改。")

    add_main_heading(doc, "5. 预期研究成果")
    add_numbered(doc, "① ", "形成一套可复现的成人 ICU DHF 操作性表型、T12 风险集和三态结局定义及其审计工件。")
    add_numbered(doc, "② ", "开发并内部验证一个与竞争风险估计目标一致的 48 h 早期血流动力学恶化或死亡预测模型。")
    add_numbered(doc, "③ ", "完成本院锁模后外部验证，明确模型校准、临床净获益、测量差异和适用范围；必要时另行报告再校准。")
    add_numbered(doc, "④ ", "完成硕士学位论文，并在结果质量允许的前提下形成学术论文、会议交流或院内静默验证方案。")

    add_main_heading(doc, "（三）研究基础")
    add_main_heading(doc, "1. 与本项目有关的研究工作积累和已取得的研究工作成绩")
    add_subheading(doc, "1.1 申请者前期工作积累")
    add_numbered(doc, "① 数据与编程基础：", "已掌握 MIMIC-IV 的患者、住院、ICU、检验、生命体征、影像、尿量和治疗相关表结构，可使用 SQL、Python 与 R 完成数据抽取、清洗、链接、建模和结果复核。")
    add_numbered(doc, "② 临床问题重构能力：", "已将旧固定二分类课题修正为由 T12 风险集、竞争事件和可观察性共同定义的预后预测问题，并形成研究逻辑链与统计分析合同。")
    add_numbered(doc, "③ 质量控制与复现能力：", "已建立变量字典、语义 allowlist、数据可用时间规则、运行登记、输入哈希、回归测试和 fail-closed 入口，能够追溯每次运行的输入、脚本和输出。")
    add_numbered(doc, "④ 双库衔接基础：", "已完成院内成人候选的初步整合、时间门控和病例复核工作台，并提前识别心超选择、报告时间和医嘱代理对外部验证的影响。")

    add_subheading(doc, "1.2 与本项目相关的前期研究结果")
    add_paragraph(doc, "截至 2026-09-20，前期工作已完成双库候选分母和数据源的首轮核对，形成 MIMIC 影像标注包、院内证据链与病例复核清单；实验室 raw 合同已在 PostgreSQL 全库和 BigQuery 固定审计快照上通过聚合质量门。上述成果证明数据路径、时间审计和清洗规则可执行，并揭示了 T12 后才可用结果、错误体液和 episode 错配等潜在偏倚来源。")
    add_paragraph(doc, "这些数字均属于候选分母、数据覆盖或质量控制。最终 DHF 人数、T12 风险集、三态结局、模型系数与性能必须来自同一冻结运行；当前 AI 预审核不是临床金标准，旧模型不是正式结果。该边界已写入项目复现合同和正式运行门控。")

    add_main_heading(doc, "2. 已具备的实验、资料等条件，尚缺少的条件和拟解决途径")
    add_subheading(doc, "2.1 已具备条件")
    add_paragraph(doc, "已具备 MIMIC-IV 3.1 数据访问和本地 PostgreSQL/BigQuery 审计条件，本院已导出的 ICU 文书、床旁心超、BNP、乳酸和目标医嘱可用于候选表型及结局可测性核对；同时具备 Git 版本控制、SQL/Python/R 运行环境、结构化运行登记和聚合质量门。")
    add_subheading(doc, "2.2 尚缺少条件")
    add_paragraph(doc, "当前尚缺双库最终 DHF 临床裁决、完整 T12 风险集与三态结局冻结；MIMIC 300 条主标注尚待临床确认，正式评阅者间一致性需要第二位标注者完成 60 条盲法复核；院内治疗执行级 eMAR/泵速资料暂不完整；完整 v3.3 依赖链和最终模型仍处于 `allow_final_run=false` 状态。")
    add_subheading(doc, "2.3 拟解决途径")
    add_numbered(doc, "① ", "优先完成对队列定义影响最大的病例级裁决和规则阴性抽样，逐例保存 A/B/C、替代解释、证据时间、来源与裁决状态。")
    add_numbered(doc, "② ", "用医嘱、护理和文书先构建可核查的治疗组成与观察完整性，明确其代理性质；eMAR 到位后按同一冻结队列重跑敏感性。")
    add_numbered(doc, "③ ", "补齐 v3.3 上游依赖与数据库硬门，只有表型、风险集、结局和特征均冻结后才开放正式模型运行。")
    add_numbered(doc, "④ ", "所有最终数字从同一冻结运行生成，预先封存阈值、折索引、随机种子、插补器、模型对象、预测值和指标定义，防止结果导向修改。")


def build_docx():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(TEMPLATE)
    update_cover(doc)
    remove_old_body(doc)
    add_body_content(doc)
    prune_unused_image_relationships(doc)
    # Metadata
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "浙江大学医学院开题报告（2026-09-20 更新）"
    doc.core_properties.author = "占舒羽"
    doc.core_properties.comments = "依据 CS_AHF_landmark24/project_control 当前正式研究资料生成；旧结局和历史可行性结果已移除。"
    doc.save(REPORT_PATH)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    write_figure_contracts()
    save_png_timeline(ASSET_DIR / "figure1_timeline" / "figure1_timeline.png")
    save_svg_timeline(ASSET_DIR / "figure1_timeline" / "figure1_timeline.svg")
    save_png_pipeline(ASSET_DIR / "figure2_pipeline" / "figure2_pipeline.png")
    save_svg_pipeline(ASSET_DIR / "figure2_pipeline" / "figure2_pipeline.svg")
    build_pptx(PPTX_PATH)
    build_docx()
    print(REPORT_PATH)
    print(PPTX_PATH)


if __name__ == "__main__":
    main()
