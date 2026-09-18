#!/usr/bin/env python3
"""Build editable Excel deliverables from the project evidence matrix."""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "literature_review" / "01_核心文献证据矩阵.csv"
TARGET = ROOT / "literature_review" / "核心文献整体分类与课题支撑_可编辑.xlsx"


def format_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for idx, column in enumerate(ws.columns, 1):
        max_len = max(len(str(c.value or "")) for c in column)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 12), 42)


def main():
    with SOURCE.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
        headers = list(rows[0])

    wb = Workbook()
    ws = wb.active
    ws.title = "核心证据矩阵"
    ws.append(headers)
    for row in rows:
        ws.append([row[h] for h in headers])
    format_sheet(ws)

    scope = wb.create_sheet("最终研究定义")
    scope.append(["字段", "冻结口径"])
    definitions = [
        ("目标人群", "成人首次ICU stay；strict AHF12；landmark前early sepsis12；0-12h无overt CS"),
        ("AHF时间边界", "ICU入科时或入科后最初12h已存在；不能笼统声称全部在入ICU前已诊断"),
        ("Landmark", "ICU intime + 12h"),
        ("预测变量窗", "ICU 0-12h"),
        ("预测窗", "ICU 12-60h"),
        ("主结局", "treatment-escalation-based hemodynamic deterioration"),
        ("主结局事件", "334/2424 (13.78%)"),
        ("严格次要结局", "mixed-shock/CS-like proxy: 133/2424 (5.49%)"),
        ("研究定位", "严格AHF人群中的early-sepsis富集主队列；early sepsis是纳入条件而非事后亚组"),
    ]
    for item in definitions:
        scope.append(item)
    format_sheet(scope)

    status = wb.create_sheet("PDF交付状态")
    status.append(["编号", "文献", "类别", "高光PDF", "可编辑", "状态/备注"])
    entries = [
        ("00", "Metra 2023", "WHF定义", "00_Metra_2023_WHF临床共识_高光.pdf", "是", "已生成"),
        ("01", "DeVore 2014", "结局与预后", "01_DeVore_2014_WHF结局与预后_高光.pdf", "是", "已生成"),
        ("02", "DeVore 2016", "静态风险模型", "02_DeVore_2016_ADHERE风险模型_高光.pdf", "是", "已生成"),
        ("03", "Rahman 2022", "AHF到CS", "03_Rahman_2022_AHF到CS早期预测_高光.pdf", "是", "已生成"),
        ("04", "Chang 2022", "跨场景CS", "04_Chang_2022_CS早期预测_高光.pdf", "是", "已生成"),
        ("05", "Zhang 2024", "HF+sepsis", "05_Zhang_2024_HF合并Sepsis死亡预测_高光.pdf", "是", "已生成"),
        ("06", "Gao 2026", "直接竞争HD", "06_Gao_2026_CVICU血流动力学恶化_高光.pdf", "是", "已生成"),
        ("07", "Beer 2024", "AHF恶化到CS", "07_Beer_2024_AHF心脏恶化到CS_高光.pdf", "是", "已生成"),
        ("08", "Hu 2024", "动态CS+外部验证", "08_Hu_2024_动态CS外部验证_高光.pdf", "是", "已生成"),
        ("09", "Essay 2020", "ICU动态特征", "", "", "Zotero条目存在；本地PDF缺失，待恢复"),
        ("10", "Greene 2023", "WHF术语", "10_Greene_2023_WHF术语框架_高光.pdf", "是", "已生成"),
        ("11", "TRIPOD+AI 2024", "报告规范", "11_Collins_2024_TRIPOD_AI报告规范_高光.pdf", "是", "已生成"),
        ("12", "PROBAST+AI 2025", "偏倚评估", "12_Moons_2025_PROBAST_AI偏倚工具_高光.pdf", "是", "已生成"),
    ]
    for entry in entries:
        status.append(entry)
    format_sheet(status)

    wb.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
