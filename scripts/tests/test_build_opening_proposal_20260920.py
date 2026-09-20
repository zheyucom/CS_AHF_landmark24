#!/usr/bin/env python3
"""Regression tests for the reproducible opening-proposal DOCX builder."""

from __future__ import annotations

import importlib.util
import re
import tempfile
import unittest
from pathlib import Path

from docx import Document


SCRIPT = Path(__file__).resolve().parents[1] / "build_opening_proposal_20260920.py"


def load_builder():
    spec = importlib.util.spec_from_file_location("opening_proposal_builder", SCRIPT)
    if spec is None or spec.loader is None:
        raise AssertionError(f"cannot import builder: {SCRIPT}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def build_document(builder):
    project_root = SCRIPT.parents[1]
    delivery = project_root / "deliverables" / "opening_proposal_20260920"
    builder.TEMPLATE = delivery / (
        "开题报告_ICU失代偿性心力衰竭患者早期血流动力学恶化或死亡预测模型_"
        "占舒羽_9.20.docx"
    )
    builder.ASSET_DIR = delivery / "figure_assets"
    with tempfile.TemporaryDirectory() as tmp:
        builder.OUT_DIR = Path(tmp)
        builder.REPORT_PATH = builder.OUT_DIR / "proposal.docx"
        builder.build_docx()
        return Document(builder.REPORT_PATH)


class CitationFormattingTests(unittest.TestCase):
    def test_numeric_citations_are_superscript_but_time_window_is_not(self):
        builder = load_builder()
        doc = Document()
        paragraph = builder.add_paragraph(
            doc,
            "独立研究结论。[3] 保留时间窗 [T0,T12)，联合方法依据。[13,14]",
        )

        superscript = [run.text for run in paragraph.runs if run.font.superscript]
        baseline = [run.text for run in paragraph.runs if not run.font.superscript]
        self.assertEqual(["[3]", "[13,14]"], superscript)
        self.assertIn("[T0,T12)", "".join(baseline))


class BuilderPathTests(unittest.TestCase):
    def test_default_output_is_a_distinct_refined_copy(self):
        builder = load_builder()

        self.assertNotEqual(builder.TEMPLATE, builder.REPORT_PATH)
        self.assertIn("_精修版", builder.REPORT_PATH.stem)
        self.assertEqual(
            builder.TEMPLATE.parent,
            builder.REPORT_PATH.parent,
        )


class OutputDocumentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.builder = load_builder()
        cls.doc = build_document(cls.builder)

    def test_reference_list_numbers_remain_on_the_baseline(self):
        reference_paragraphs = [
            paragraph
            for paragraph in self.doc.paragraphs
            if re.match(r"^\[\d+\]", paragraph.text)
        ]
        self.assertEqual(18, len(reference_paragraphs))
        for paragraph in reference_paragraphs:
            marker = re.match(r"^\[\d+\]", paragraph.text).group(0)
            marker_runs = [run for run in paragraph.runs if run.text.startswith(marker)]
            self.assertEqual(1, len(marker_runs), paragraph.text)
            self.assertFalse(marker_runs[0].font.superscript, paragraph.text)

    def test_hu_beer_and_rahman_are_cited_as_independent_claims(self):
        status_start = next(
            i for i, paragraph in enumerate(self.doc.paragraphs)
            if paragraph.text == "2.2 国内外研究现状"
        )
        reference_start = next(
            i for i, paragraph in enumerate(self.doc.paragraphs)
            if paragraph.text == "3. 主要参考文献"
        )
        status_paragraphs = self.doc.paragraphs[status_start + 1:reference_start]
        expected = {"Hu 等": "[3]", "Beer 等": "[4]", "Rahman 等": "[5]"}
        for author, citation in expected.items():
            matches = [paragraph for paragraph in status_paragraphs if author in paragraph.text]
            self.assertEqual(1, len(matches), author)
            paragraph = matches[0]
            mentioned_authors = [name for name in expected if name in paragraph.text]
            self.assertEqual([author], mentioned_authors, paragraph.text)
            superscript = [run.text for run in paragraph.runs if run.font.superscript]
            self.assertIn(citation, superscript, paragraph.text)

    def test_every_in_text_numeric_citation_is_a_superscript_run(self):
        for paragraph in self.doc.paragraphs:
            if re.match(r"^\[\d+\]", paragraph.text):
                continue
            for match in self.builder.CITATION_TOKEN_RE.finditer(paragraph.text):
                token = match.group(0)
                token_runs = [run for run in paragraph.runs if token in run.text]
                self.assertTrue(token_runs, paragraph.text)
                self.assertTrue(
                    all(run.font.superscript for run in token_runs),
                    paragraph.text,
                )

    def test_progress_snapshot_matches_the_current_project_state(self):
        text_parts = [paragraph.text for paragraph in self.doc.paragraphs]
        for table in self.doc.tables:
            for row in table.rows:
                text_parts.extend(cell.text for cell in row.cells)
        text = "\n".join(text_parts)
        for expected in (
            "29 项实验室特征来源矩阵",
            "063A v3",
            "120/120",
            "181/181",
            "0 blocker",
            "正式队列、三态结局和模型尚未冻结",
        ):
            self.assertIn(expected, text)
        self.assertNotIn("91/91", text)

    def test_main_basis_heading_appears_only_once(self):
        matches = [
            paragraph
            for paragraph in self.doc.paragraphs
            if paragraph.text == "（一）选题依据"
        ]
        self.assertEqual(1, len(matches))


if __name__ == "__main__":
    unittest.main(verbosity=2)
